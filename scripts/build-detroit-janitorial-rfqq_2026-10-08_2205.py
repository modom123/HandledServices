# ============================================================================
# FILE    : scripts/build-detroit-janitorial-rfqq_2026-10-08_2205.py
# PROJECT : Handled (HandledServices) — AI-run home & business services
# CREATED : 2026-10-08_2205 UTC
# PURPOSE : Response package for City of Detroit RFQQ 184795 — Detroit Supply Schedule 2 (DSS #2),
#           Janitorial/Custodial Services. Built from the City's RFQQ, Rider #1, Rider #2 and
#           Attachments A, B and D (copies in docs/RFQQ184795_CITY_FORMS/).
#             docs/HANDLED_RFQQ184795_TECHNICAL_PROPOSAL_2026-10-08_2205.docx (+ .pdf preview)
#                 Rider #1 item (3): cover letter, org chart, resumes of key personnel, equipment, references,
#                 laid out to answer Attachment B and the RFQQ's Technical Factors 1–3.
#             docs/HANDLED_RFQQ184795_ATTACHMENT_A_FILLED_2026-10-08_2205.docx
#                 The City's Attachment A questionnaire with Handled's known answers filled in.
#             docs/HANDLED_RFQQ184795_SAMPLE_EMPLOYMENT_APPLICATION_2026-10-08_2205.docx (+ .pdf)
#                 Attachment D item 4 — fair-chance application (no criminal-history questions).
#           Anything in [[double brackets]] prints highlighted yellow: fill or confirm before submitting.
#           Supersedes scripts/build-detroit-janitorial-rfqq_2026-10-08_2154.py (written before the RFQQ was readable).
#           Run: python3 scripts/build-detroit-janitorial-rfqq_2026-10-08_2205.py
# ============================================================================
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

STAMP = "2026-10-08_2205"
FORMS = "docs/RFQQ184795_CITY_FORMS"
OUT_TP = f"docs/HANDLED_RFQQ184795_TECHNICAL_PROPOSAL_{STAMP}.docx"
OUT_A = f"docs/HANDLED_RFQQ184795_ATTACHMENT_A_FILLED_{STAMP}.docx"
OUT_APP = f"docs/HANDLED_RFQQ184795_SAMPLE_EMPLOYMENT_APPLICATION_{STAMP}.docx"
LOGO = "apps/web/public/brand/handled-lockup.png"
CO = "Handled Services LLC"
PHONE, EMAIL, SITE = "(313) 639-9373", "info@handledsvc.com", "handledsvc.com"
RFQ = "RFQQ 184795"
TITLE = "Detroit Supply Schedule #2 — Janitorial/Custodial Services"
ADDR = "[[Street address]], Detroit, MI [[ZIP]]"
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
INK = RGBColor(0x1A, 0x1F, 0x1C)
SOFT = RGBColor(0x5B, 0x66, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ── helpers ──────────────────────────────────────────────────────────────────
def runs(p, text, bold=False, size=None, color=None, italic=False):
    """Add text to a paragraph; [[...]] segments become yellow-highlighted fill-in fields."""
    for part in re.split(r"(\[\[.*?\]\])", text):
        if not part:
            continue
        fill = part.startswith("[[")
        r = p.add_run(f"[{part[2:-2]}]" if fill else part)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if fill: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill); tcPr.append(sh)


def set_cell(cell, text, **kw):
    cell.text = ""
    lines = text.split("\n")
    runs(cell.paragraphs[0], lines[0], **kw)
    for ln in lines[1:]:
        runs(cell.add_paragraph(), ln, **kw)


class Doc:
    def __init__(self, header_text):
        d = self.d = Document()
        sec = d.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin, sec.bottom_margin = Inches(1.2), Inches(0.9)
        sec.header_distance, sec.footer_distance = Inches(0.4), Inches(0.35)
        n = d.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5); n.paragraph_format.space_after = Pt(5)
        for name, size in (("Heading 1", 15), ("Heading 2", 12)):
            s = d.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
            s.font.color.rgb = GREEN if name == "Heading 1" else INK
            s.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8); s.paragraph_format.space_after = Pt(4)
        hdr = sec.header
        ht = hdr.add_table(rows=1, cols=2, width=Inches(6.7)); ht.autofit = False
        l, r = ht.rows[0].cells; l.width, r.width = Inches(3.2), Inches(3.5)
        l.paragraphs[0].add_run().add_picture(LOGO, width=Inches(1.9))
        rp = r.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        runs(rp, header_text, size=8.5, color=SOFT)
        hdr.paragraphs[0].text = ""
        rule = hdr.add_paragraph(); pPr = rule._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "10", "w:space": "1", "w:color": "1F7A4D"}.items(): b.set(qn(k), v)
        bdr.append(b); pPr.append(bdr)
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runs(fp, f"{CO}  ·  {RFQ}  ·  {PHONE}  ·  {EMAIL}", size=7.5, color=SOFT)

    def para(self, text="", bold=False, size=None, color=None, align=None, after=None, italic=False):
        p = self.d.add_paragraph(); runs(p, text, bold, size, color, italic)
        if align is not None: p.alignment = align
        if after is not None: p.paragraph_format.space_after = Pt(after)
        return p

    def bullets(self, items):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            if isinstance(it, tuple): runs(p, it[0], bold=True); runs(p, " " + it[1])
            else: runs(p, it)

    def h1(self, t): self.d.add_heading(t, level=1)
    def h2(self, t): self.d.add_heading(t, level=2)

    def table(self, header, rows, widths, label_col=False):
        t = self.d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; set_cell(c, h, bold=True, size=9.5, color=WHITE); shade(c, "1F7A4D")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row): set_cell(cells[i], v, size=9.5, bold=(label_col and i == 0))
        t.autofit = False
        for i, w in enumerate(widths): t.columns[i].width = Inches(w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def page_break(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path, title, pdf=True):
        self.d.core_properties.title = title; self.d.core_properties.author = CO
        self.d.core_properties.comments = f"Generated {STAMP} UTC by scripts/build-detroit-janitorial-rfqq_{STAMP}.py"
        self.d.save(path); print("wrote", path)
        if pdf:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", "docs", path], check=False,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            print("wrote", path.replace(".docx", ".pdf"))


# ═════════════════════════════════════════════════════════════════════════════
# 1. TECHNICAL PROPOSAL
# ═════════════════════════════════════════════════════════════════════════════
D = Doc(f"{RFQ}\n{TITLE}")
C = WD_ALIGN_PARAGRAPH.CENTER
D.d.sections[0].different_first_page_header_footer = True

# ── cover
D.para(after=36)
cp = D.d.add_paragraph(); cp.alignment = C; cp.add_run().add_picture(LOGO, width=Inches(3.6))
D.para(after=26)
D.para("TECHNICAL PROPOSAL / STATEMENT OF QUALIFICATIONS", bold=True, size=12, color=SOFT, align=C, after=4)
D.para("Janitorial/Custodial Services", bold=True, size=24, color=INK, align=C, after=4)
D.para("City of Detroit · Detroit Supply Schedule #2 (DSS #2)", size=12, color=SOFT, align=C, after=2)
D.para(f"Request for Qualifications {RFQ}", bold=True, size=12, color=GREEN, align=C, after=6)
D.para("Items offered: ITEM 1 Janitorial/Custodial Services · ITEM 3 Garbage/Trash Removal Service"
       " · [[ITEM 2 Cleaning Services, Steam and Pressure — keep only if equipment and 5-year experience are met]]",
       size=10, align=C, after=40)
D.para("Submitted to", size=9.5, color=SOFT, align=C, after=0)
D.para("City of Detroit, Office of the Chief Financial Officer — Office of Contracting and Procurement", bold=True, align=C, after=0)
D.para("Coleman A. Young Municipal Center · 2 Woodward Avenue, Suite 1008 · Detroit, MI 48226", size=9.5, align=C, after=22)
D.para("Submitted by", size=9.5, color=SOFT, align=C, after=0)
D.para(CO + " — Certified Detroit Business [[certification type and #]]", bold=True, size=12, align=C, after=0)
D.para(ADDR, size=9.5, align=C, after=0)
D.para(f"{PHONE} · {EMAIL} · {SITE}", size=9.5, align=C, after=0)
D.para("Submission date: [[October __, 2026]] · Quarterly review: [[Q_ 2026/2027]]", size=9.5, align=C)
D.page_break()

# ── contents
D.h1("Contents")
D.table(["Section", "Responds to"], [
    ["1. Cover Letter / Proposal Introduction", "Rider #1 (3) cover letter · Attachment B-A"],
    ["2. Items Offered", "RFQQ Part 1 — Schedule Item Numbers"],
    ["3. Factor 1 — Municipal Experience", "Offer Preparation (h) Factor 1 (2-page limit)"],
    ["4. Factor 2 — Past Performance and References", "Factor 2 · Rider #1 (3) references · Attachment A Part 1 §3"],
    ["5. Factor 3 — Quality Control", "Factor 3 (2-page limit)"],
    ["6. Solution / Approach to the Statement of Work", "Attachment B-B · RFQQ Scope of Work"],
    ["7. Response Timeline / Schedule", "Attachment B-C"],
    ["8. Technical Approach — Procedures", "Attachment B-D"],
    ["9. Organization Chart", "Rider #1 (3) · Attachment A §5a"],
    ["10. Key Personnel and Resumes", "Rider #1 (3) · Attachment A §5b–c · Position Descriptions"],
    ["11. Equipment", "Rider #1 (3) · ITEM 2 equipment qualifications"],
    ["12. Insurance", "RFQQ Insurance section"],
    ["13. Licenses, Registrations and Compliance", "Attachment B-A · Attachment D · Consolidated Affidavits"],
    ["14. Submission Checklist", "Rider #1 (d) · Attachment D"],
], [3.4, 3.3])
D.page_break()

# ── 1. cover letter
D.h1("1. Cover Letter / Proposal Introduction")
D.para("[[October __, 2026]]", after=8)
D.para("City of Detroit — Office of Contracting and Procurement\nColeman A. Young Municipal Center\n2 Woodward Avenue, Suite 1008\n"
       "Detroit, MI 48226\nAttn: [[Arnita C______]], Buyer", after=8)
D.para(f"Re: {RFQ} — {TITLE}", bold=True, after=8)
D.para("Dear [[Ms. ______]]:")
D.para(f"{CO} (“Handled”) submits this Statement of Qualifications in response to {RFQ}, Detroit Supply Schedule #2 for "
       "Janitorial/Custodial Services, and requests prequalification for ITEM 1 (Janitorial/Custodial Services) and "
       "ITEM 3 (Garbage/Trash Removal Service)[[, and ITEM 2 (Cleaning Services, Steam and Pressure)]].")
D.para("Handled is a Detroit-headquartered, [[certified Detroit-Based Business (DBB) / Detroit-Headquartered Business (DHB) / "
       "Detroit Small Business (DSB)]] that delivers commercial cleaning, floor care, trash and debris removal, window and "
       "pressure cleaning through supervised, vetted crews drawn from Detroit neighborhoods. Every service visit is recorded "
       "against a checklist with time-stamped before/after photos, so ordering departments can see what was done, when and by whom.")
D.para("Handled commits to perform all work ordered under this Schedule in accordance with the requirements of the RFQQ, including "
       "the Scope of Work task frequencies, the position description qualifications for each labor category, the insurance "
       "requirements, and all City of Detroit ordinances and affidavits. By this submission Handled takes no exception to the "
       "terms and conditions of the RFQQ. Handled acknowledges Rider #1 and Rider #2[[ and Amendment(s) No. __]].")
D.para("Licenses and registrations held: Michigan limited liability company in good standing ([[MI LARA ID #]]); "
       "City of Detroit Oracle supplier registration ([[Supplier #]]); Detroit business certification ([[type / # / expiration]]); "
       "[[City of Detroit business license, if required]]; [[SAM.gov UEI, if applicable]]. Copies are attached.")
D.para("Our single point of contact for this RFQQ is [[Name, Title]], at "
       f"{PHONE} or {EMAIL}. Thank you for your consideration.")
D.para("Sincerely,", after=24)
D.para(f"[[Name]]\n[[Title — must be listed on the LLC Certificate of Authority]], {CO}")
D.page_break()

# ── 2. items
D.h1("2. Items Offered")
D.table(["ITEM", "Description (RFQQ Part 1)", "Offered", "Basis"], [
    ["1", "Janitorial/Custodial Services — interior and exterior building cleaning; periodic heavy cleaning; daily inspections; "
          "bio-hazard clean-up and disposal per OSHA as needed", "Yes",
     "Supervised crews, floor-care and carpet extraction equipment; bio-hazard through OSHA-trained staff [[or licensed partner]]"],
    ["2", "Cleaning Services, Steam and Pressure — DDOT Gilbert & Shoemaker Terminal bays (6×/yr) and Rosa Parks Transit Center (18×/yr)",
     "[[Yes / No]]", "[[Only if the required sweeper-scrubbers, 2,000-gal recovery tank, 2,000–4,000 psi water blaster are owned/leased "
                     "AND 5 years of high-pressure cleaning at similar facilities can be shown — directly or through a named subcontractor]]"],
    ["3", "Garbage/Trash Removal Service — City locations including parks, recreation facilities and outside City buildings", "Yes",
     "Collection, bagging, hauling to approved pick-up/disposal points; [[vehicles and disposal arrangements]]"],
], [0.5, 2.6, 0.7, 2.9])
D.para("Proposed labor categories (fully loaded rates to be provided per the RFQQ pricing instructions): Janitorial/Custodial "
       "Supervisor; Janitor/Custodian; [[Trash Removal Driver/Laborer]]; [[Pressure Washing Technician]]. Each complies with the "
       "RFQQ position descriptions (Supervisor: 5 years supervisory experience, HS diploma/GED; Custodian: 2 years experience, "
       "HS diploma/GED; both able to lift 30 lb and speak and understand English).", size=9.5)

# ── 3. factor 1
D.h1("3. Factor 1 — Municipal Experience")
D.para("(RFQQ limit: two pages. Minimum of two (2) years of municipal experience required.)", size=9, color=SOFT, italic=True)
D.h2("3.1 Years of municipal experience")
D.para("[[State the exact number of years Handled — or, for a joint venture, each JV party — has provided janitorial/custodial or "
       "trash removal services to a municipality or other public agency, and list the agencies. The RFQQ requires at least two "
       "years. If Handled itself does not yet have two years, this response must come from a joint-venture partner or the "
       "submission should wait; do not count non-municipal work here.]]")
D.h2("3.2 Employees and resources")
D.table(["Measure", "Response"], [
    ["Total workforce available to this Schedule", "[[__]] (employees [[__]] · subcontracted crew members [[__]])"],
    ["Detroit-resident workers", "[[__]] ([[__ %]])"],
    ["Supervisors meeting the 5-year requirement", "[[__]]"],
    ["Custodians meeting the 2-year requirement", "[[__]]"],
    ["Surge capacity (additional vetted workers within 72 hours)", "[[__]]"],
    ["Vehicles / major equipment", "See Section 11"],
], [3.3, 3.4], label_col=True)
D.h2("3.3 Organizational and accounting controls")
D.bullets([
    ("Order control.", "Every City purchase order is entered as a job with its PO number, ordering department, site, scope, and "
                       "frequency; no work is dispatched without a PO and no invoice is issued without a matching PO line."),
    ("Time and service records.", "Crews clock in and out by site with GPS-stamped records; each visit closes only after the "
                                  "checklist and photos are complete. These records back every invoice."),
    ("Accounting.", "Books kept in Xero, reconciled daily against payment and payout records; separate job-cost tracking by PO "
                    "and ordering department; [[outside CPA: firm name]] reviews [[annually]]."),
    ("Payroll and compliance.", "Workers paid weekly; certificates of insurance, background checks, and training records tracked "
                                "with expiration alerts — lapsed items automatically block assignment."),
    ("Separation of duties.", "Dispatch, quality inspection, and billing are performed by different people; [[owner]] approves "
                              "all invoices to the City."),
])
D.h2("3.4 Intended use of subcontractors")
D.para("[[Choose one and delete the other:]]")
D.para("[[(a) Handled will perform all work with its own employees.]]")
D.para("[[(b) Handled will use the following subcontractor(s): legal name, address, services supplied, % of work. Each will sign "
       "the Covenant of Equal Opportunity and meet the same insurance, vetting and training requirements, and Handled remains "
       "responsible for all work. This must match Attachment A §4 and any subcontracting plan.]]")

# ── 4. factor 2
D.h1("4. Factor 2 — Past Performance and References")
D.para("Three relevant current or past contracts within the past three (3) years, for Handled and any proposed subcontractors. "
       "[[Only list contracts Handled or a named subcontractor actually performed, with the client's permission. If fewer than "
       "three exist, keep those you have and use the statement in 4.4.]]", after=6)
for n in (1, 2, 3):
    D.h2(f"4.{n} Contract {n}")
    D.table(["Field", "Response"], [
        ["(a) Description of the project", "[[Scope, facility type, square footage, frequency, crew size]]"],
        ["(b) Contract number", "[[ ]]"],
        ["(c) Contract / order amount", "[[$ ]]"],
        ["(d) Name of organization", "[[ ]]"],
        ["(e) Point of contact", "[[Name, title, address, phone, email]]"],
        ["(f) Current status", "[[Completed / in progress — start date, completion date]]"],
        ["Performed by", "[[Handled / subcontractor name]] · key personnel: [[ ]]"],
    ], [2.3, 4.4], label_col=True)
D.h2("4.4 Statement if past performance is insufficient")
D.para("[[Use only if applicable:]] “Handled Services LLC affirmatively states that it possesses insufficient relevant past "
       "performance that is directly related or similar to the efforts required by the Janitorial/Custodial Services, beyond the "
       "contracts listed above.”", italic=True)

# ── 5. factor 3
D.h1("5. Factor 3 — Quality Control")
D.para("(RFQQ limit: two pages.)", size=9, color=SOFT, italic=True)
D.h2("5.1 Individuals responsible for quality control")
D.table(["Role", "Name", "QC responsibility"], [
    ["Quality Control Manager", "[[Name]]", "Owns the QC program; monthly unannounced inspections; corrective action plans; QC reporting to each ordering department"],
    ["Account Manager", "[[Name]]", "Single point of contact for City departments; receives complaints and urgent requests; tracks open issues to closure"],
    ["Site Supervisors", "[[Names]]", "Shift inspections against the site checklist; sign off every visit; first response to deficiencies"],
], [1.6, 1.3, 3.8])
D.h2("5.2 How potential problems and solutions are handled")
D.bullets([
    ("Prevent.", "Each site has a written site plan built from the RFQQ task frequencies (daily through annual) and the department's "
                 "instructions; crews are trained on it before starting."),
    ("Detect.", "Every visit: room-by-room checklist plus time-stamped before/after photos, reviewed before the visit is closed. "
                "Supervisors inspect each shift; the QC Manager performs unannounced inspections at least [[monthly]] per site using "
                "a 100-point form. Hazardous conditions (burned-out lights, loose railings, ceiling tiles, exposed wiring, broken "
                "windows) are reported verbally to the supervisor and in writing to the Project Manager with the date observed, as the RFQQ requires."),
    ("Correct.", "A deficiency reported by a department is acknowledged within [[1]] hour and re-cleaned within [[24]] hours at no cost. "
                 "Inspection scores below [[90]] trigger a written corrective action plan delivered to the Project Manager within [[3]] business days."),
    ("Prevent recurrence.", "Root cause recorded (training, staffing, supplies, equipment); repeated deficiencies remove a worker or crew from City sites."),
])
D.h2("5.3 Quality when responding to urgent requirements")
D.bullets([
    "Emergency line answered 24/7 at [[phone]]; a supervisor is on call every night and weekend.",
    "Standby crew of [[__]] vetted workers who already hold City-site clearance, so urgent work is never staffed by untrained people.",
    "Emergency kits (spill response, PPE, bio-hazard bags, wet vacuums, signage) staged in [[__]] vehicles.",
    "Urgent jobs use the same checklist-and-photo close-out as routine work, and the supervisor inspects before release.",
])
D.h2("5.4 Quality control across multiple purchase orders and ordering activities")
D.bullets([
    "Each PO is its own job record with its own site plan, schedule, checklist, assigned supervisor, and department contact — so "
    "requirements from DDOT, GSD, Parks & Recreation or other departments never mix.",
    "A weekly operations review covers every open PO: visits completed vs. scheduled, inspection scores, open deficiencies, staffing.",
    "Each ordering department receives a [[monthly]] QC report for its POs: on-time visits, inspection scores, deficiencies and time to correct, emergency response times.",
    "Capacity check before accepting any new PO: Handled will not accept work that would pull supervisors or crews from existing City sites.",
])

# ── 6. solution / approach
D.h1("6. Solution / Approach to the Statement of Work")
D.para("Handled will perform the specified cleaning and janitorial services at all City office facilities and work sites in a purchase "
       "order, at the frequencies in the RFQQ Scope of Work, which we treat as the minimum standard. Cyclical starting dates for "
       "monthly, quarterly, semi-annual and annual tasks will be set with the City Project Manager and Plant/Site Manager during the "
       "first month, and any change to duties, procedures or materials will be made only with the Project Manager's prior written approval.")
D.h2("6.1 ITEM 1 — Janitorial/Custodial Services: task frequencies")
D.table(["Frequency", "Tasks (per RFQQ Scope of Work)"], [
    ["Daily", "Dusting (no visible dust or cobwebs, treated cloths/HEPA tools); damp wiping; thorough sweeping; lobby glass spot-cleaning; "
              "restroom cleaning in the RFQQ sequence with closure signage and an agreed closing schedule; showers/locker rooms; "
              "lunchroom/kitchen; stainless steel with the grain; drinking fountains; empty waste receptacles; rubbish removal to dumpsters/compactors; "
              "HEPA vacuuming of carpets; dust/damp mopping; wet mopping and power scrubbing of heavy-traffic areas; restroom floors auto-scrubbed with enzymes"],
    ["Every other day", "Spray-buff hard floors"],
    ["Weekly", "Spot-clean walls, doors, glass, display cases; remove graffiti, tape and expired notices; clean whiteboards; eaves and exterior "
               "graffiti; ash receptacles and entry areas; remove/replace and extract carpet runners; scrub restroom and stairwell floors and walls; "
               "buff and wax (e.g., GSD, DDOT)"],
    ["Monthly", "Interior office glass and partitions; door hardware, switches, push plates, kick plates; air bars, vents, louvers and diffusers; "
                "kitchen appliances inside and out, vent hoods"],
    ["Quarterly", "Locker exteriors; piping and mechanical equipment surfaces; strip and wax hard floors (min. 25% solids finish); "
                  "carpet extraction (commercial truck-mounted equipment)"],
    ["Semi-annually", "Storage areas, basements, ramps; fans, A/C and vents; interior and exterior windows up to 12 ft"],
    ["Annually", "Lighting diffusers, lenses and reflectors up to 12 ft; pressure-wash exteriors of buildings two stories or less; "
                 "wash interior walls, steam-clean concrete walls"],
    ["Immediately", "Emergency clean-up assigned by the Project Manager; grease, oil and other spills; hazardous-condition reporting"],
], [1.3, 5.4], label_col=True)
D.para("Crew conduct: workers will not clean or tamper with computer terminals, screens or control panels, disturb papers, open desk "
       "drawers or cabinets, or use telephones, printers or other City office equipment.", size=9.5)
D.h2("6.2 Bio-hazard clean-up (ITEM 1, as needed)")
D.para("Bodily fluids and waste, solid waste, and chemical clean-ups or spills will be handled under OSHA 29 CFR 1910.1030 (Bloodborne "
       "Pathogens) and 1910.1200 (Hazard Communication): trained staff with hepatitis B vaccination offered, PPE, EPA-registered "
       "disinfectants at label dwell times, red-bag segregation, and disposal through [[licensed medical-waste hauler]]. "
       "[[If an outside licensed firm will do this work, name it in Section 3.4.]]")
D.h2("6.3 ITEM 2 — Steam and pressure cleaning [[delete if not offered]]")
D.para("Six cleanings per year at the Gilbert Terminal bays (~138,800 sq ft) and Shoemaker Terminal bays (~155,800 sq ft), and eighteen per "
       "year at Rosa Parks Transit Center (~65,000 sq ft), plus as-needed cleanings after spills. Method: trash and debris removal, "
       "scarifying, high-pressure hot-water washing with spin-jet attachments, and EPA-determined non-hazardous chemicals to remove oil, "
       "grease and scum; all dirty water collected into recovery tanks and disposed of under federal, state and local law, with "
       "manifests at no additional cost. Equipment: see Section 11.")
D.h2("6.4 ITEM 3 — Garbage/trash removal service")
D.para("Collection and removal of garbage and trash at City locations including parks, recreation facilities and outside City buildings, "
       "on the route and frequency in each PO: litter pick-up, emptying and relining receptacles, bagging, and delivery to the approved "
       "pick-up point or hauling to [[licensed disposal/transfer facility]]. Vehicles: [[number and type]]. Route sheets with "
       "GPS-stamped photos at each stop are available to the ordering department.")
D.h2("6.5 Products and safety")
D.bullets([
    "Green Seal / EPA Safer Choice products where performance allows; EPA-registered disinfectants for restrooms and high-touch surfaces.",
    "Safety Data Sheets on site; secondary containers labeled; color-coded microfiber to prevent cross-contamination; HEPA vacuums emptied daily.",
    "Wet-floor signage; incidents reported to the Project Manager within [[24]] hours.",
    "Key and access control: keys signed out and logged, codes never written on keys, lost keys reported immediately.",
])

# ── 7. timeline
D.h1("7. Response Timeline / Schedule")
D.table(["Event", "Handled commitment"], [
    ["Quote in response to a department RFQ under the Schedule", "Within [[3]] business days (site walk-through scheduled within [[2]] business days if needed)"],
    ["Start of recurring service after PO", "Within [[5–10]] business days"],
    ["One-time / periodic service (floor strip, carpet extraction, windows)", "Scheduled within [[5]] business days"],
    ["Urgent same-day request", "Crew on site the same day for requests received by [[12:00 p.m.]]"],
    ["Emergency (spill, flood, bio-hazard), 24/7", "Phone acknowledgement within [[30]] minutes; crew on site within [[4]] hours"],
    ["Deficiency reported by department", "Corrected within [[24]] hours at no charge"],
    ["Hazardous condition observed", "Verbal report immediately; written report to Project Manager same day"],
], [3.2, 3.5], label_col=True)

# ── 8. technical approach
D.h1("8. Technical Approach — Procedures")
D.para("How Handled identifies, evaluates and communicates throughout a purchase order:")
D.table(["Step", "Procedure", "Communication to the City"], [
    ["1. Identify", "Receive RFQ/PO; walk the site with the department contact; measure areas; note special conditions (security, hours, sensitive rooms)", "Written walk-through notes"],
    ["2. Evaluate", "Build site plan from RFQQ frequencies; staff to the labor categories; confirm equipment and supplies; quote at Schedule rates", "Quote and site plan for approval"],
    ["3. Start", "Orientation of crew to site rules; key and access issue; first-week supervisor present every shift", "Start-up confirmation; contact list"],
    ["4. Perform", "Checklist and photo record each visit; periodic tasks on the agreed cycle", "Visit records available on request"],
    ["5. Inspect", "Supervisor and QC Manager inspections; deficiencies corrected within [[24]] hours", "Monthly QC report"],
    ["6. Invoice", "Monthly invoice per PO with service dates, hours by labor category and PO number", "Invoice with backup records"],
    ["7. Review", "Quarterly review meeting with each department on recurring POs", "Meeting notes and action items"],
], [1.0, 3.6, 2.1])

# ── 9. org chart
D.h1("9. Organization Chart")
oc = D.d.add_table(rows=4, cols=3); oc.style = "Table Grid"; oc.alignment = WD_TABLE_ALIGNMENT.CENTER
top = oc.cell(0, 0).merge(oc.cell(0, 2)); set_cell(top, "[[Name]] — Owner / Managing Member\nContract executive; signs on behalf of the LLC", bold=True, size=9.5, color=WHITE); shade(top, "1F7A4D")
for i, t in enumerate(["[[Name]] — Account Manager\nCity point of contact; POs, scheduling, invoicing",
                       "[[Name]] — Quality Control Manager\nInspections, corrective action, safety & training",
                       "[[Name]] — Operations / Dispatch\nStaffing, routes, equipment, emergency line"]):
    set_cell(oc.cell(1, i), t, bold=True, size=9); shade(oc.cell(1, i), "E3EFE8")
for i, t in enumerate(["Janitorial/Custodial Supervisors ([[__]])\n5+ yrs supervisory", "Janitors/Custodians ([[__]])\n2+ yrs experience",
                       "Trash removal crews ([[__]])\n[[Pressure-wash technicians ([[__]])]]"]):
    set_cell(oc.cell(2, i), t, size=9)
bot = oc.cell(3, 0).merge(oc.cell(3, 2)); set_cell(bot, "[[Subcontractor(s), if any — name and services]]", size=9)
D.para()

# ── 10. key personnel
D.h1("10. Key Personnel and Resumes")
D.para("Resumes for each key person are attached. Summary:", after=6)
D.table(["Name / role", "Years relevant experience", "Qualifications", "Detroit resident"], [
    ["[[Name]] — Account Manager", "[[__]]", "[[ ]]", "[[Y/N]]"],
    ["[[Name]] — Quality Control Manager", "[[__]]", "[[OSHA 10/30, BBP trainer, ISSA CIMS, etc.]]", "[[Y/N]]"],
    ["[[Name]] — Janitorial/Custodial Supervisor", "[[5+]] supervisory", "HS diploma/GED; [[ ]]", "[[Y/N]]"],
    ["[[Name]] — Janitorial/Custodial Supervisor", "[[5+]] supervisory", "HS diploma/GED; [[ ]]", "[[Y/N]]"],
], [2.3, 1.3, 2.2, 0.9])
D.para("Number of employees that can be dedicated to these services: [[__]]. Resume format for each attachment: name, role on this "
       "Schedule, employment history with dates and employers, janitorial/custodial or supervisory duties, certifications, education.", size=9.5)

# ── 11. equipment
D.h1("11. Equipment")
D.table(["Equipment", "Qty", "Owned / leased / subcontractor"], [
    ["Auto scrubber (walk-behind)", "[[ ]]", "[[ ]]"],
    ["Low-speed floor machine / rotary buffer", "[[ ]]", "[[ ]]"],
    ["High-speed burnisher", "[[ ]]", "[[ ]]"],
    ["Wet/dry vacuums", "[[ ]]", "[[ ]]"],
    ["HEPA backpack and upright vacuums", "[[ ]]", "[[ ]]"],
    ["Carpet extraction — truck-mounted (preferred) or commercial portable", "[[ ]]", "[[ ]]"],
    ["Pressure washer (hot water)", "[[ ]]", "[[ ]]"],
    ["Ladders (up to 12 ft work height), signage, spill and bio-hazard kits", "[[ ]]", "[[ ]]"],
    ["Service vehicles / trash removal trucks", "[[ ]]", "[[ ]]"],
    ["ITEM 2 only — Tennant 8410 sweeper-scrubber (or approved equal), one per terminal", "[[2]]", "[[ ]]"],
    ["ITEM 2 only — 2,000-gal poly tank for dirty water recovery, one per terminal", "[[2]]", "[[ ]]"],
    ["ITEM 2 only — 2,000–4,000 psi water blaster with spin-jet floor attachment, one per terminal", "[[2]]", "[[ ]]"],
    ["ITEM 2 only — Advance CS 7000 sweeper-scrubber with power washer (or approved equal), Rosa Parks", "[[1]]", "[[ ]]"],
], [4.2, 0.6, 1.9])

# ── 12. insurance
D.h1("12. Insurance")
D.para("Handled will carry, at its own expense, the insurance the RFQQ requires, with the City of Detroit named as additional insured on "
       "the Commercial General Liability policy, and 30 days' prior notice of cancellation or reduction. Certificates will be provided "
       "before any work begins.", after=6)
D.table(["Coverage (RFQQ requirement)", "Minimum limit", "Carrier / status"], [
    ["Workers' Compensation", "Michigan statutory", "[[Carrier · in force / bound on award]]"],
    ["Employer's Liability", "$500,000 each accident · $500,000 each disease · $500,000 policy limit", "[[ ]]"],
    ["Commercial General Liability (City as additional insured)", "$1,000,000 per occurrence · $2,000,000 aggregate", "[[ ]]"],
    ["Automobile Liability — owned, hired and non-owned; Michigan No-Fault PIP/PPI", "$1,000,000 combined single limit", "[[ ]]"],
    ["MCS-90 endorsement (only if hazardous waste is transported)", "Per RFQQ", "[[N/A or carrier]]"],
], [2.8, 2.4, 1.5])

# ── 13. compliance
D.h1("13. Licenses, Registrations and Compliance")
D.bullets([
    ("Detroit certification.", "This RFQQ is set aside for certified Detroit businesses. Handled's certification: [[type / # / expiration — attach]]."),
    ("Fair-chance hiring.", "Handled's employment application does not ask about criminal history; convictions are not inquired into "
                            "or considered until after an interview or a determination that the applicant is qualified, per Detroit City "
                            "Code §§ 17-5-261 through 17-5-266. Sample application attached."),
    ("Equal opportunity.", "Handled and all subcontractors will sign and comply with the Covenant of Equal Opportunity; subcontractor "
                           "covenants will go to the Human Rights Department before work starts."),
    ("Clearances.", "Income tax and accounts receivable clearances [[requested on __ / received]] (bit.ly/detroitclearances); evidence attached."),
    ("Financial responsibility.", "Financial statements for the previous three years (balance sheet, income statement, profit and loss) "
                                  "attached[[, with an explanation of any negative results or of years before the company was formed]]."),
])

# ── 14. checklist
D.h1("14. Submission Checklist")
D.para("Rider #1 replaced the submission list: upload through Oracle (1) tax clearances and affidavits, (2) three years of financial "
       "statements, and (3) this technical proposal with cover letter, organization chart, key-personnel resumes, equipment and "
       "references. Attachment D lists the forms.", after=6)
D.table(["Document", "Source", "Notary", "Done"], [
    ["Technical Proposal (this document) with resumes attached", "Rider #1 (3)", "", "☐"],
    ["Attachment A — Respondent Questionnaire (completed)", "City form", "", "☐"],
    ["Attachment B — Proposal Introduction/Approach (answered by Sections 1, 6, 7, 8)", "City form", "", "☐"],
    ["Certificate of Authority — LLC version only", "Attachment D #1", "", "☐"],
    ["Amendment Form — Rider #1, Rider #2 and any later amendments", "Attachment D #2", "", "☐"],
    ["Consolidated Affidavits (Hiring Policy, Slavery Era/Prison/Detention, Covenant of Equal Opportunity, Political Contributions — enter NONE if none)", "Attachment D #3", "Yes", "☐"],
    ["Exhibit D — Non-Collusion Affidavit", "City form", "Yes", "☐"],
    ["Sample Employment Application", "Attachment D #4", "", "☐"],
    ["Income & Revenue Tax Clearance — evidence of request", "Attachment D #5", "", "☐"],
    ["Three (3) years financial statements (not tax returns)", "Attachment D #6 · Rider #1 (2)", "", "☐"],
    ["SAM.gov registration screenshot (only if grant-funded)", "Attachment D #7", "", "☐"],
    ["Detroit business certification", "Set-aside", "", "☐"],
    ["Certificates of insurance (or broker letter)", "Insurance section", "", "☐"],
], [3.9, 1.5, 0.6, 0.7])
D.save(OUT_TP, f"{CO} — {RFQ} Technical Proposal")


# ═════════════════════════════════════════════════════════════════════════════
# 2. ATTACHMENT A — fill the City's own form
# ═════════════════════════════════════════════════════════════════════════════
A = Document(f"{FORMS}/Attachment_A_Respondent_Questionnaire_BLANK.docx")


def fill(table, answers):
    for row in table.rows:
        label = row.cells[0].text.strip().lower()
        for key, val in answers.items():
            if label.startswith(key.lower()):
                c = row.cells[-1]; set_cell(c, val, size=10)
                break


poc = {"Name": "[[Name]]", "Title": "[[Title]]", "Address": ADDR, "E-mail": EMAIL, "Phone": PHONE}
fill(A.tables[0], poc)
fill(A.tables[1], {"Name": "[[Name]]", "Title": "[[Account Manager / Owner]]", "Address": ADDR, "E-mail": "[[email]]", "Phone": PHONE})
fill(A.tables[2], {
    "Full legal business name": CO,
    "Full legal business address": ADDR + "\nPlace of performance: City of Detroit facilities as ordered",
    "Business entity": "Limited Liability Company (LLC)",
    "Current tax status": "[[Tax classification, e.g. single-member LLC / partnership / S-corp]] · FEIN [[XX-XXXXXXX]]",
    "State company formed": "Michigan",
    "Company phone": PHONE,
    "Website": SITE,
    "Number of years in business": "[[__]]",
    "Average number of employees": "[[Year 1: __ · Year 2: __ · Year 3: __]]",
    "Does your company have experience": "[[No / Yes — list City of Detroit contracts from the last 5 years]]",
    "Identify any claims": "[[None]]",
    "Identify any projects": "[[None]]",
})
ref_blank = {"Name of Reference": "[[ ]]", "Project Name": "[[ ]]", "Client Location": "[[ ]]", "Contact Person Name": "[[ ]]",
             "Contact Person Title": "[[ ]]", "Contact Person Phone": "[[ ]]", "Contact Person E-mail": "[[ ]]",
             "Dates of Service": "[[mm/yy – mm/yy]]", "Description of Services": "[[ ]]", "Identify respondent": "[[ ]]"}
for i in (3, 4, 5, 7): fill(A.tables[i], ref_blank)
fill(A.tables[6], {"Do you intend": "[[No / Yes — legal name, address, services supplied]]"})
for p in A.paragraphs:  # the City's template says "Moving and Relocation" — keep their text, flag it
    if "Moving and Relocation" in p.text:
        r = p.add_run("  [Note: answered for Janitorial/Custodial Services, DSS #2.]"); r.font.size = Pt(9); r.italic = True
p = A.add_paragraph()
runs(p, "Part 2 §5 response: organization chart — Technical Proposal Section 9; resumes — attached to the Technical Proposal; "
        "employees that can be dedicated to these services: [[__]].", size=10)
A.core_properties.comments = f"Filled {STAMP} UTC by scripts/build-detroit-janitorial-rfqq_{STAMP}.py"
A.save(OUT_A); print("wrote", OUT_A)


# ═════════════════════════════════════════════════════════════════════════════
# 3. SAMPLE EMPLOYMENT APPLICATION (Attachment D #4) — fair-chance
# ═════════════════════════════════════════════════════════════════════════════
E = Doc(f"Employment Application\nJanitorial / Custodial")
E.para("Employment Application — Janitorial/Custodial Positions", bold=True, size=16, color=INK, after=2)
E.para(f"{CO} · {ADDR} · {PHONE} · {EMAIL}", size=9, color=SOFT, after=8)
E.para(f"{CO} is an equal opportunity employer. We do not ask about criminal history on this application. Consistent with "
       "Detroit City Code §§ 17-5-261 through 17-5-266, we will not inquire into or consider criminal convictions until after "
       "an interview or a determination that you are qualified for the position. Any later background check is individually "
       "assessed and you will have the opportunity to respond.", size=9.5, italic=True)
E.h2("Position")
E.table(["Field", ""], [["Position applied for", "☐ Janitor/Custodian   ☐ Janitorial/Custodial Supervisor   ☐ Trash Removal   ☐ Other: ____"],
                        ["Available start date", ""], ["Shifts available", "☐ Early morning   ☐ Day   ☐ Evening   ☐ Overnight   ☐ Weekends"],
                        ["Full-time / part-time", "☐ Full-time   ☐ Part-time   ☐ On-call"]], [2.2, 4.5], label_col=True)
E.h2("Applicant information")
E.table(["Field", ""], [["Full name", ""], ["Address", ""], ["City / State / ZIP", ""], ["Phone", ""], ["Email", ""],
                        ["Are you legally authorized to work in the United States?", "☐ Yes   ☐ No"],
                        ["Are you at least 18 years old?", "☐ Yes   ☐ No"],
                        ["Can you lift and move objects up to 30 lb, with or without reasonable accommodation?", "☐ Yes   ☐ No"],
                        ["Can you speak and understand English?", "☐ Yes   ☐ No"]], [3.6, 3.1], label_col=True)
E.h2("Education")
E.table(["Field", ""], [["High school diploma or GED?", "☐ Yes   ☐ No"], ["Other training / certifications (OSHA, BBP, floor care, etc.)", ""]],
        [3.6, 3.1], label_col=True)
E.h2("Work experience (most recent first)")
for n in (1, 2, 3):
    E.table([f"Employer {n}", ""], [["Employer / location", ""], ["Dates (mm/yy – mm/yy)", ""], ["Job title", ""],
                                    ["Duties (cleaning, floor care, supervision, etc.)", ""], ["Supervisor / phone — may we contact?", "☐ Yes   ☐ No"]],
            [2.6, 4.1], label_col=True)
E.table(["Field", ""], [["Total years of janitorial/custodial experience", ""], ["Years supervising janitorial/custodial staff", ""]],
        [3.6, 3.1], label_col=True)
E.h2("References (not related to you)")
E.table(["Name", "Relationship", "Phone"], [["", "", ""], ["", "", ""], ["", "", ""]], [2.6, 2.3, 1.8])
E.h2("Applicant certification")
E.para("I certify that the information on this application is true and complete. I understand that false or misleading information "
       "may result in my application being rejected.", size=9.5, after=18)
E.para("Signature ______________________________________    Date ________________")
E.save(OUT_APP, f"{CO} — Employment Application (fair-chance)")
