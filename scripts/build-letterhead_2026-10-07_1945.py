# ============================================================================
# FILE    : scripts/build-letterhead_2026-10-07_1945.py
# PROJECT : Handled (HandledServices) — AI-run home & business services
# CREATED : 2026-10-07_1945 UTC
# PURPOSE : Official Handled letterhead, two files:
#             docs/HANDLED_LETTERHEAD_2026-10-07_1945.pdf   — blank page to print on or attach behind a letter
#             docs/HANDLED_LETTERHEAD_2026-10-07_1945.docx  — Word letterhead: logo + contact block in the header,
#                                                            footer line, ready to type a letter (Word or Google Docs)
#           Run: python3 scripts/build-letterhead_2026-10-07_1945.py
# ============================================================================
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as pdfcanvas
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = "apps/web/public/brand/handled-lockup.png"
PHONE, EMAIL, SITE = "(313) 639-9373", "info@handledsvc.com", "handledsvc.com"
FOOT = "Home & business services in Michigan and Washington  ·  Vetted, background-checked, insured pros  ·  30-day make-it-right guarantee"
GREEN, INK, SOFT = colors.HexColor("#1f7a4d"), colors.HexColor("#1a1f1c"), colors.HexColor("#5b6660")

# ── PDF
c = pdfcanvas.Canvas("docs/HANDLED_LETTERHEAD_2026-10-07_1945.pdf", pagesize=letter)
c.setTitle("Handled letterhead"); c.setAuthor("Handled Services LLC")
c.drawImage(LOGO, 0.75 * inch, 10.05 * inch, width=2.6 * inch, height=2.6 * inch * 240 / 891, mask="auto")
c.setFillColor(INK); c.setFont("Helvetica-Bold", 9); c.drawRightString(7.75 * inch, 10.5 * inch, "Handled Services LLC")
c.setFont("Helvetica", 8.5); c.setFillColor(SOFT)
c.drawRightString(7.75 * inch, 10.35 * inch, f"{PHONE}  ·  {EMAIL}"); c.drawRightString(7.75 * inch, 10.2 * inch, SITE)
c.setStrokeColor(GREEN); c.setLineWidth(1.4); c.line(0.75 * inch, 9.92 * inch, 7.75 * inch, 9.92 * inch)
c.setStrokeColor(colors.HexColor("#dfe5e1")); c.setLineWidth(0.5); c.line(0.75 * inch, 0.68 * inch, 7.75 * inch, 0.68 * inch)
c.setFillColor(SOFT); c.setFont("Helvetica", 7.5); c.drawCentredString(4.25 * inch, 0.5 * inch, FOOT)
c.showPage(); c.save()

# ── Word
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.75); sec.top_margin = Inches(1.35); sec.bottom_margin = Inches(0.9)
sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.35)
style = d.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(11)

hdr = sec.header
t = hdr.add_table(rows=1, cols=2, width=Inches(7))
t.autofit = False
left, right = t.rows[0].cells
left.width, right.width = Inches(3.6), Inches(3.4)
left.paragraphs[0].add_run().add_picture(LOGO, width=Inches(2.6))
rp = right.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = rp.add_run("Handled Services LLC\n"); r.bold = True; r.font.size = Pt(9)
r = rp.add_run(f"{PHONE}  ·  {EMAIL}\n{SITE}"); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x5B, 0x66, 0x60)
hdr.paragraphs[0].text = ""
# green rule under the header
line = hdr.add_paragraph()
pPr = line._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
for k, v in {"w:val": "single", "w:sz": "12", "w:space": "1", "w:color": "1F7A4D"}.items(): bottom.set(qn(k), v)
bdr.append(bottom); pPr.append(bdr)

fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(FOOT); fr.font.size = Pt(7.5); fr.font.color.rgb = RGBColor(0x5B, 0x66, 0x60)

for txt in ["[Date]", "", "[Name]\n[Title]\n[Company]\n[Street address]\n[City, State ZIP]", "", "Dear [Name],", "",
            "[Type your letter here.]", "", "Sincerely,", "", "", "[Your name]\n[Title], Handled Services LLC\n" + f"{PHONE} · {EMAIL}"]:
    d.add_paragraph(txt)
d.core_properties.title = "Handled letterhead"; d.core_properties.author = "Handled Services LLC"
d.save("docs/HANDLED_LETTERHEAD_2026-10-07_1945.docx")
print("wrote letterhead PDF + DOCX")
