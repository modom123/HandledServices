# ============================================================================
# FILE    : scripts/build-detroit-janitorial-rfqq_2026-10-08_2154.py
# PROJECT : Handled (HandledServices) — AI-run home & business services
# CREATED : 2026-10-08_2154 UTC
# PURPOSE : Statement of Qualifications for the City of Detroit RFQQ 184795 —
#           Detroit Supply Schedule 2: Janitorial/Custodial Services.
#             docs/HANDLED_RFQQ184795_DETROIT_JANITORIAL_SOQ_2026-10-08_2154.docx — editable, fill the yellow fields
#             docs/HANDLED_RFQQ184795_DETROIT_JANITORIAL_SOQ_2026-10-08_2154.pdf  — preview (LibreOffice)
#           Anything in [[double brackets]] prints highlighted yellow: it must be filled or confirmed
#           against the RFQQ, Rider #1 (item numbers) and Rider #2 (annual schedule) before submitting.
#           Run: python3 scripts/build-detroit-janitorial-rfqq_2026-10-08_2154.py
# ============================================================================
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

STAMP = "2026-10-08_2154"
OUT = f"docs/HANDLED_RFQQ184795_DETROIT_JANITORIAL_SOQ_{STAMP}.docx"
LOGO = "apps/web/public/brand/handled-lockup.png"
CO = "Handled Services LLC"
PHONE, EMAIL, SITE = "(313) 639-9373", "info@handledsvc.com", "handledsvc.com"
RFQ = "RFQQ 184795"
TITLE = "Detroit Supply Schedule 2 — Janitorial/Custodial Services"
FOOT = f"{CO}  ·  {RFQ} Statement of Qualifications  ·  {PHONE}  ·  {EMAIL}"
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
INK = RGBColor(0x1A, 0x1F, 0x1C)
SOFT = RGBColor(0x5B, 0x66, 0x60)

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin, sec.bottom_margin = Inches(1.2), Inches(0.9)
sec.header_distance, sec.footer_distance = Inches(0.4), Inches(0.35)
sec.different_first_page_header_footer = True
normal = d.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
for name, size in (("Heading 1", 15), ("Heading 2", 12)):
    s = d.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = GREEN if name == "Heading 1" else INK
    s.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8); s.paragraph_format.space_after = Pt(4)


def runs(p, text, bold=False, size=None, color=None):
    """Add text to a paragraph; [[...]] segments become yellow-highlighted fill-in fields."""
    for i, part in enumerate(re.split(r"(\[\[.*?\]\])", text)):
        if not part:
            continue
        fill = part.startswith("[[")
        r = p.add_run(part[2:-2] if fill else part)
        r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if fill:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.text = f"[{r.text}]"
    return p


def para(text="", bold=False, size=None, color=None, align=None, after=None):
    p = d.add_paragraph()
    runs(p, text, bold, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def bullets(items):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            runs(p, it[0], bold=True); runs(p, " " + it[1])
        else:
            runs(p, it)


def h1(t): d.add_heading(t, level=1)
def h2(t): d.add_heading(t, level=2)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_fill); tcPr.append(sh)


def table(header, rows, widths):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; runs(c.paragraphs[0], h, bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF)); shade(c, "1F7A4D")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; runs(cells[i].paragraphs[0], v, size=9.5, bold=(i == 0 and len(header) == 2))
    t.autofit = False
    for i, w in enumerate(widths): t.columns[i].width = Inches(w)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── running header / footer (pages 2+)
hdr = sec.header
ht = hdr.add_table(rows=1, cols=2, width=Inches(6.7)); ht.autofit = False
l, r = ht.rows[0].cells; l.width, r.width = Inches(3.2), Inches(3.5)
l.paragraphs[0].add_run().add_picture(LOGO, width=Inches(1.9))
rp = r.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
runs(rp, f"{RFQ}\n{TITLE}", size=8.5, color=SOFT)
hdr.paragraphs[0].text = ""
rule = hdr.add_paragraph(); pPr = rule._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
for k, v in {"w:val": "single", "w:sz": "10", "w:space": "1", "w:color": "1F7A4D"}.items(): b.set(qn(k), v)
bdr.append(b); pPr.append(bdr)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
runs(fp, FOOT, size=7.5, color=SOFT)

# ── cover page
d.add_paragraph().paragraph_format.space_after = Pt(40)
cp = d.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.add_run().add_picture(LOGO, width=Inches(3.6))
para("", after=30)
para("STATEMENT OF QUALIFICATIONS", bold=True, size=13, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("Janitorial/Custodial Services", bold=True, size=24, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para(f"City of Detroit · Detroit Supply Schedule Program · Schedule 2", size=12, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(f"Request for Qualifications {RFQ}", bold=True, size=12, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=50)
para("Submitted to", size=9.5, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("City of Detroit, Office of Contracting and Procurement", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Coleman A. Young Municipal Center · 2 Woodward Avenue · Detroit, MI 48226", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Oracle Supplier Portal · procurementInthecloud@detroitmi.gov · arnitac@detroitmi.gov", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
para("Submitted by", size=9.5, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(CO, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("[[Street address]] · Detroit, MI [[ZIP]]", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(f"{PHONE} · {EMAIL} · {SITE}", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Submission date: [[October __, 2026]]", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

# ── transmittal letter
para("[[October __, 2026]]", after=10)
para("City of Detroit, Office of Contracting and Procurement\nColeman A. Young Municipal Center\n2 Woodward Avenue, Suite 1008\nDetroit, MI 48226\nAttn: [[Arnita C______, Buyer]]", after=10)
para(f"Re: {RFQ} — {TITLE}", bold=True, after=10)
para("Dear [[Ms. ______]]:")
para(f"{CO} (“Handled”) is pleased to submit this Statement of Qualifications to be placed on the Detroit Supply Schedule "
     "for Janitorial/Custodial Services. We are a Detroit-based company that delivers commercial and facility cleaning through a "
     "vetted network of local, insured cleaning professionals, coordinated by a modern dispatch, inspection and invoicing platform "
     "and accountable to one point of contact.")
para("We are offering qualifications for the Schedule 2 items identified in Section 3 of this submission. For each, the City "
     "and participating municipal jurisdictions will get:")
bullets([
    ("One accountable vendor.", "A dedicated account manager, a single ordering channel and one consolidated invoice per task order."),
    ("Detroit people doing Detroit work.", "Our crews are recruited from Detroit neighborhoods; every worker is ID-verified and background-checked before their first shift."),
    ("Proof the work was done.", "Time-stamped before/after photos and an inspection checklist for every service visit, available to the requesting agency."),
    ("Fast response.", "Scheduled, on-call and emergency response, with a make-it-right commitment: any deficiency is re-cleaned at no cost."),
])
para("We understand this is an indefinite-delivery, indefinite-quantity schedule, that no minimum volume is guaranteed, and that "
     "orders will be placed by individual agencies through task orders at the rates established under the schedule. We agree to "
     "the terms of the RFQQ and its attachments, riders and addenda [[list any addenda received, or “none”]].")
para("Our authorized contact for this submission is below. Thank you for your consideration.")
para("Sincerely,", after=26)
para(f"[[Name]]\n[[Title]], {CO}\n{PHONE} · {EMAIL}")
page_break()

# ── 1. offeror information
h1("1. Offeror Information")
table(["Item", "Response"], [
    ["Legal business name", CO],
    ["DBA", "Handled · HandledServices"],
    ["Business address", "[[Street address]], Detroit, MI [[ZIP]]"],
    ["Mailing address (if different)", "[[Same as above]]"],
    ["Primary contact / authorized signer", "[[Name, Title]]"],
    ["Phone / email", f"{PHONE} · {EMAIL}"],
    ["Website", SITE],
    ["Business structure / state", "Limited Liability Company · Michigan"],
    ["Federal EIN", "[[XX-XXXXXXX]]"],
    ["City of Detroit Oracle supplier number", "[[Supplier #]]"],
    ["Year established", "[[2026]]"],
    ["NAICS codes", "561720 Janitorial Services · 561790 Other Services to Buildings and Dwellings · 562111 Solid Waste Collection (trash/debris removal)"],
    ["Detroit business certifications", "[[Detroit-Based Business (DBB) / Detroit Small Business (DSB) / Detroit-Resident Business (DRB) — pending or certificate #]]"],
    ["Other certifications", "[[MBE / WBE / SBE / none]]"],
    ["Schedule 2 items offered", "See Section 3"],
], [2.3, 4.4])

# ── 2. company overview
h1("2. Company Overview")
para(f"{CO} is a Detroit services company serving commercial and residential customers across Michigan, with operations "
     "also in Washington State. Our core service lines are commercial and residential cleaning, window and carpet cleaning, "
     "power washing, trash and debris removal, grounds and snow work, and handyman services.")
para("We operate differently from a traditional janitorial contractor. Instead of a large fixed payroll, we maintain a "
     "deep bench of vetted, insured local cleaning professionals and small crews, and our operations platform matches each work "
     "order to the qualified, nearest available crew. That model gives the City three things a fixed-staff vendor struggles to:")
bullets([
    ("Surge capacity.", "Elections, events, weather emergencies and building moves can be staffed without waiting on new hires."),
    ("Coverage citywide.", "Crews are matched by proximity, so sites from the riverfront to Eight Mile get a crew that is close by."),
    ("Measured quality.", "Every visit is inspected by photo and checklist and every crew is rated; consistently high-rated crews are assigned first and under-performers are removed from City work."),
])
para("A human operations team supervises every account. Our platform handles scheduling, routing, inspection records and "
     "invoicing; it never replaces the site supervisor or the account manager the City talks to.")

# ── 3. items offered
h1("3. Schedule 2 Items Offered")
para("Handled offers qualifications for the items marked “Yes” below. Item numbers correspond to Rider #1 (DSS #2 items). "
     "[[Confirm every item number and description against Rider #1; delete rows for items not offered.]]", after=6)
table(["Item No.", "Service", "Offered", "Typical deliverables"], [
    ["[[2-__]]", "Routine janitorial / custodial cleaning (offices, public buildings)", "Yes", "Daily/weekly cleaning per agency frequency schedule; trash and recycling; dusting; restrooms; floors"],
    ["[[2-__]]", "Restroom sanitation and restocking", "Yes", "Disinfection of all fixtures, restock of paper and soap, odor control"],
    ["[[2-__]]", "Floor care — sweep, mop, strip, wax, burnish", "Yes", "VCT strip and refinish, scheduled burnishing, entry-mat care"],
    ["[[2-__]]", "Carpet and upholstery cleaning", "Yes", "Hot-water extraction, spot treatment, encapsulation between deep cleans"],
    ["[[2-__]]", "Interior and exterior window cleaning", "Yes", "Ground-level and reachable-height glass; partitions; entry doors"],
    ["[[2-__]]", "Disinfection / sanitization services", "Yes", "EPA List N disinfectants, high-touch point program, outbreak response cleaning"],
    ["[[2-__]]", "Day porter / matron services", "Yes", "On-site attendant during business hours; lobby, restroom and spill response"],
    ["[[2-__]]", "Post-construction and move-in/move-out cleaning", "Yes", "Rough, final and touch-up cleans; debris removal"],
    ["[[2-__]]", "Special event setup and cleanup", "Yes", "Pre-event cleaning, during-event porters, post-event cleanup and trash removal"],
    ["[[2-__]]", "Power washing (sidewalks, entries, dumpster pads)", "Yes", "Hot/cold pressure washing; gum and graffiti removal where specified"],
    ["[[2-__]]", "Trash, debris and bulk item removal", "Yes", "Hauling to licensed disposal; recycling where available"],
    ["[[2-__]]", "Emergency / on-call cleaning", "Yes", "Water intrusion cleanup, bio-hazard area isolation (see note), after-hours response"],
], [0.75, 2.1, 0.65, 3.2])
para("Note: blood and body-fluid cleanup beyond routine spill response, mold remediation and asbestos-related work will be "
     "performed only by appropriately licensed partners where the task order requires it, and identified to the City before work begins.", size=9, color=SOFT)

# ── 4. technical approach
h1("4. Technical Approach and Service Standards")
h2("4.1 Starting a task order")
bullets([
    ("Request.", "The agency issues a task order or request through its preferred channel (email, phone, or our client portal)."),
    ("Walk-through.", "For recurring or large work, we schedule a site walk-through with the agency contact within [[3]] business days, at no charge."),
    ("Scope and quote.", "Within [[2]] business days of the walk-through, we deliver a written scope: tasks, frequencies, square footage, hours and the price at schedule rates."),
    ("Site plan.", "On approval, we issue a site plan: crew assignment, supervisor, schedule, access/key control procedure, products, and the inspection checklist."),
    ("Start.", "Service begins on the agency’s date, typically within [[5–10]] business days of approval for recurring service."),
])
h2("4.2 Standard cleaning frequencies (adjusted per site)")
table(["Area / task", "Daily", "Weekly", "Monthly / periodic"], [
    ["Trash & recycling removal, liner replacement", "✓", "", ""],
    ["Restrooms: disinfect fixtures, mirrors, partitions; restock", "✓", "", ""],
    ["High-touch disinfection (handles, rails, switches, elevator buttons)", "✓", "", ""],
    ["Hard floors: sweep and damp mop; carpet: vacuum traffic areas", "✓", "", ""],
    ["Dust horizontal surfaces, desks (cleared), sills", "", "✓", ""],
    ["Full vacuum including edges; spot-clean carpet", "", "✓", ""],
    ["Clean interior glass and entry doors", "", "✓", ""],
    ["High dusting, vents, baseboards", "", "", "Monthly"],
    ["Burnish hard floors", "", "", "Monthly / as scheduled"],
    ["Strip and refinish VCT; carpet extraction", "", "", "Quarterly–semi-annual"],
    ["Exterior windows, power wash entries", "", "", "Semi-annual / as ordered"],
], [3.5, 0.7, 0.7, 1.8])
h2("4.3 Response times")
table(["Request type", "Response commitment"], [
    ["Emergency (spill, water intrusion, sanitation hazard)", "Acknowledge within 30 minutes; crew on site within [[4]] hours, 24/7"],
    ["Urgent (same-day need, e.g. event, inspection)", "Crew on site same day if requested by [[12:00 p.m.]]"],
    ["Routine one-time request", "Scheduled within [[3]] business days"],
    ["Recurring service start", "Within [[5–10]] business days of task order approval"],
    ["Deficiency reported by agency", "Corrected within [[24]] hours at no charge"],
], [3.0, 3.7])

# ── 5. staffing
h1("5. Staffing, Vetting and Training")
para("Every person who enters a City facility on our behalf has passed the same vetting:")
bullets([
    "Government photo ID verification and a criminal background check before the first assignment; [[re-checked annually]]. Additional agency-specific clearance (e.g., police, water, airport, or courthouse facilities) completed as the agency requires.",
    "Proof of general liability insurance (and workers’ compensation for crews with employees) collected and tracked for expiration; lapsed coverage automatically blocks assignment.",
    "Signed confidentiality, key-control and code-of-conduct agreements.",
    "Training before City work: OSHA Hazard Communication (chemical labels and SDS), bloodborne pathogens awareness, slip/trip/fall prevention, proper PPE, and our site checklist and photo-inspection procedure.",
    "Uniformed and badged on site, with a supervisor reachable by phone for every shift.",
])
table(["Role", "Responsibility", "Name"], [
    ["Account manager", "Single point of contact for the City; task orders, scheduling, invoicing, issue resolution", "[[Name]]"],
    ["Operations / quality manager", "Inspections, crew performance, corrective action, safety", "[[Name]]"],
    ["Site supervisors", "On-shift supervision, checklist sign-off, key control", "[[Assigned per site]]"],
    ["Cleaning crews", "Service delivery per site plan", "Vetted local crews"],
], [1.6, 3.7, 1.4])

# ── 6. quality control
h1("6. Quality Control Plan")
bullets([
    ("Checklist for every visit.", "Each site plan includes a room-by-room checklist the crew completes and signs digitally."),
    ("Before/after photos.", "Time- and location-stamped photos of key areas on every visit, reviewed before the visit is closed. Agencies can request the photo record at any time."),
    ("Supervisor inspections.", "Unannounced supervisor inspections at least [[monthly]] per site, scored on a standard 100-point form; scores below [[90]] trigger a corrective action plan."),
    ("Agency feedback.", "A one-click rating after each service and a quarterly review meeting with each agency contact on recurring accounts."),
    ("Make-it-right.", "Any deficiency reported by the agency is re-cleaned within [[24]] hours at no charge."),
    ("Crew accountability.", "Crew ratings combine agency feedback and our inspections; repeated deficiencies remove a crew from City assignments."),
])
para("Performance indicators reported to each agency [[quarterly]]: on-time arrival rate, inspection score, deficiencies "
     "reported and time to correct, emergency response time, and invoice accuracy.")

# ── 7. safety & products
h1("7. Safety, Products and Environmental Practices")
bullets([
    "Green Seal / EPA Safer Choice certified cleaning products by default; EPA-registered disinfectants (List N where applicable) used per label dwell time.",
    "Safety Data Sheets on site for every product; secondary containers labeled per OSHA HazCom.",
    "Color-coded microfiber system to prevent cross-contamination between restrooms and other areas; HEPA-filter vacuums.",
    "Wet-floor signage and spill control on every shift; incidents reported to the agency within [[24]] hours.",
    "Recycling and waste separation according to each building’s program.",
    "Key and alarm-code control: keys logged and signed out, codes never written on keys, and lost keys reported immediately.",
])

# ── 8. insurance
h1("8. Insurance, Licenses and Compliance")
para("Handled will maintain, at a minimum, the insurance required by the RFQQ, name the City of Detroit as additional insured, "
     "and provide certificates before the first task order. [[Confirm limits against the RFQQ insurance section.]]", after=6)
table(["Coverage", "Limit", "Status"], [
    ["Commercial general liability", "[[$1,000,000 per occurrence / $2,000,000 aggregate]]", "[[In force / bound on award]]"],
    ["Workers’ compensation / employer’s liability", "Statutory / [[$500,000]]", "[[In force / bound on award]]"],
    ["Commercial auto (owned / hired & non-owned)", "[[$1,000,000 combined single limit]]", "[[In force / bound on award]]"],
    ["Janitorial bond / third-party crime", "[[$______]]", "[[In force / bound on award]]"],
    ["Umbrella / excess liability", "[[$______]]", "[[If required]]"],
], [2.6, 2.5, 1.6])
para("Handled will comply with all federal, state and City requirements that apply to this schedule, including the "
     "City’s clearance requirements (income tax, property tax and water/sewer), Human Rights Department requirements, "
     "applicable wage and hiring ordinances, and the affidavits included with this submission.")

# ── 9. local
h1("9. Detroit Business Participation and Local Hiring")
para("Handled is headquartered in Detroit and answers the phone with a 313 number. Our crews are recruited from Detroit "
     "neighborhoods through local job boards, community organizations and workforce partners, and we give priority for City "
     "assignments to Detroit-resident workers. [[State current % of workers who are Detroit residents, and Detroit certification status.]]")
bullets([
    "Partnering with [[Detroit at Work / Detroit Employment Solutions Corporation]] to recruit and train Detroit residents for City sites.",
    "Purchasing supplies from Detroit-based vendors where price and availability allow.",
    "Fast payment to our crews (weekly), which keeps reliable local workers on City accounts.",
])

# ── 10. pricing
h1("10. Pricing")
para("Handled will provide pricing in the format required by the RFQQ and Rider #2 (Annual Schedule). Rates will hold "
     "for the schedule term except as the RFQQ allows. [[Enter rates from your pricing sheet, or attach the City’s pricing form and "
     "reference it here.]]", after=6)
table(["Rate type", "Unit", "Rate"], [
    ["Routine janitorial — standard hours (Mon–Fri, 6 a.m.–6 p.m.)", "Per labor hour", "[[$____]]"],
    ["Routine janitorial — after hours / weekends", "Per labor hour", "[[$____]]"],
    ["Recurring building cleaning", "Per sq. ft. per month", "[[$____]]"],
    ["Day porter", "Per hour", "[[$____]]"],
    ["Floor strip and refinish (VCT)", "Per sq. ft.", "[[$____]]"],
    ["Floor burnish", "Per sq. ft.", "[[$____]]"],
    ["Carpet extraction", "Per sq. ft.", "[[$____]]"],
    ["Window cleaning", "Per pane / per hour", "[[$____]]"],
    ["Power washing", "Per sq. ft.", "[[$____]]"],
    ["Emergency call-out (24/7)", "Per hour, [[2]]-hour minimum", "[[$____]]"],
    ["Supplies (paper, liners, soap) when furnished by Handled", "Cost plus", "[[__ %]]"],
], [3.4, 1.8, 1.5])
para("Invoicing: one invoice per task order per month (or per one-time job), referencing the agency’s PO number, with "
     "service dates, hours or square footage, and the inspection record. Payment terms per the City’s standard terms "
     "[[Net 30]].")

# ── 11. references
h1("11. Experience and References")
para("[[Provide three references for comparable commercial cleaning work — client name, contact, phone/email, address, "
     "scope, square footage, frequency and dates. If the company is newer than the RFQQ’s experience requirement, list the "
     "principals’ and lead supervisors’ prior janitorial experience and contracts here, and state it plainly. Do not list "
     "a reference without that client’s permission.]]", after=6)
for n in (1, 2, 3):
    table(["Reference " + str(n), ""], [
        ["Organization", "[[ ]]"], ["Contact name / title", "[[ ]]"], ["Phone / email", "[[ ]]"],
        ["Facility type, size and frequency", "[[ ]]"], ["Period of performance", "[[ ]]"], ["Annual value", "[[ ]]"],
    ], [2.3, 4.4])

# ── 12. required forms
h1("12. Required Forms and Attachments")
para("The following City documents are completed, signed and included with this submission:", after=6)
table(["Document", "Included"], [
    ["Attachment A — Detroit Supply Schedule [[form title]]", "[[✓]]"],
    ["Attachment B — Detroit Supply Schedule [[form title]]", "[[✓]]"],
    ["Attachment D — Detroit Supply Schedule [[form title]]", "[[✓]]"],
    ["Rider #1 — DSS #2 Items (items offered marked)", "[[✓]]"],
    ["Rider #2 — Annual Schedule", "[[✓]]"],
    ["Consolidated Affidavits (signed and notarized as required)", "[[✓]]"],
    ["Exhibit D — Non-Collusion Affidavit (notarized)", "[[✓]]"],
    ["Combined Certificates of Authority / clearances", "[[✓]]"],
    ["Certificate of insurance (or broker letter of insurability)", "[[✓]]"],
    ["Detroit business certification (DBB/DSB/DRB), if held", "[[✓ / N/A]]"],
    ["Michigan LLC certificate of good standing; W-9", "[[✓]]"],
], [5.4, 1.3])

# ── signature
h1("13. Certification and Signature")
para(f"I certify that the information in this Statement of Qualifications is true and complete, and that I am authorized "
     f"to submit it on behalf of {CO}.", after=24)
table(["", ""], [
    ["Signature", "\n\n"], ["Name / title", "[[ ]]"], ["Date", "[[ ]]"],
], [2.3, 4.4])

d.core_properties.title = f"{CO} — {RFQ} Janitorial/Custodial Services Statement of Qualifications"
d.core_properties.author = CO
d.core_properties.comments = f"Generated {STAMP} UTC by scripts/build-detroit-janitorial-rfqq_{STAMP}.py"
d.save(OUT)
print("wrote", OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", "docs", OUT], check=False,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print("wrote", OUT.replace(".docx", ".pdf"))
